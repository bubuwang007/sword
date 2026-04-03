"""Word 文档表格封装模块."""

from __future__ import annotations

from typing import Any, Iterator

from docx.table import Table
from docx.document import Document as DocxDocument

from sword.cell import WordCell


class WordTable:
    """封装 Word 文档中的表格，提供便捷的表格操作."""

    def __init__(
        self,
        document: DocxDocument,
        rows: int,
        cols: int,
        style: str | None = None,
    ) -> None:
        """
        初始化表格.

        Args:
            document: 底层的 python-docx Document 对象.
            rows: 表格行数.
            cols: 表格列数.
            style: 表格样式（如 "Table Grid"）。
        """
        self._doc = document
        self._table = document.add_table(rows=rows, cols=cols)
        if style is not None:
            self._table.style = style

    @property
    def table(self) -> Table:
        """获取底层的 python-docx Table 对象."""
        return self._table

    @property
    def rows(self) -> int:
        """获取表格行数."""
        return len(self._table.rows)

    @property
    def cols(self) -> int:
        """获取表格列数."""
        return len(self._table.columns)

    def cell(self, row: int, col: int) -> WordCell:
        """
        获取单元格.

        Args:
            row: 行索引（从 0 开始）。
            col: 列索引（从 0 开始）。

        Returns:
            WordCell 对象。
        """
        return WordCell(self._table.cell(row, col))

    def set_style(self, style: str | None) -> None:
        """
        设置表格样式.

        Args:
            style: 样式名称（None 清除样式）。
        """
        if style is None:
            self._table.style = None
        else:
            self._table.style = style

    def iter_cells(self) -> Iterator[WordCell]:
        """
        遍历表格中所有单元格.

        Yields:
            WordCell 对象，按行优先顺序。
        """
        for row in self._table.rows:
            for cell in row.cells:
                yield WordCell(cell)

    def iter_rows(self) -> Iterator[list[WordCell]]:
        """
        按行遍历表格.

        Yields:
            每行的 WordCell 对象列表。
        """
        for row in self._table.rows:
            yield [WordCell(cell) for cell in row.cells]

    def iter_cols(self) -> Iterator[list[WordCell]]:
        """
        按列遍历表格.

        Yields:
            每列的 WordCell 对象列表。
        """
        for column in self._table.columns:
            yield [WordCell(cell) for cell in column.cells]

    def get_row(self, row_index: int) -> list[WordCell]:
        """
        获取指定行的所有单元格.

        Args:
            row_index: 行索引（从 0 开始）。

        Returns:
            WordCell 对象列表。
        """
        row = self._table.rows[row_index]
        return [WordCell(cell) for cell in row.cells]

    def get_column(self, col_index: int) -> list[WordCell]:
        """
        获取指定列的所有单元格.

        Args:
            col_index: 列索引（从 0 开始）。

        Returns:
            WordCell 对象列表。
        """
        column = self._table.columns[col_index]
        return [WordCell(cell) for cell in column.cells]

    def __enter__(self) -> WordTable:
        return self

    def __exit__(self, *args: Any) -> None:
        pass
